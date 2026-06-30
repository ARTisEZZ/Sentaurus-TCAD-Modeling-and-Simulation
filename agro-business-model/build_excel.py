#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор Excel-книги финмодели сельхозбизнеса (Томск).

Создаёт agro_business_model.xlsx с листами:
  Допущения, CAPEX, P&L, Юнит-экономика, Денежный поток (база),
  Прогноз (3 сценария), Чувствительность, Источники.

Листы CAPEX / P&L / Юнит-экономика / Денежный поток — НА ФОРМУЛАХ,
ссылающихся на лист «Допущения»: меняешь вход — пересчитывается всё.
Листы «Прогноз» и «Чувствительность» — расчётные снимки (значения).

Запуск: python3 build_excel.py   (нужен пакет openpyxl)
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import model as M

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "agro_business_model.xlsx")

# --- стили ---
H1 = Font(bold=True, size=14, color="FFFFFF")
H2 = Font(bold=True, size=11, color="FFFFFF")
BOLD = Font(bold=True)
INPUT_FONT = Font(color="0000CC")
TITLE_FILL = PatternFill("solid", fgColor="2F5496")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="FCE4D6")
GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
RUB = '# ##0 \\₽'
PCT = '0.0%'
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header_row(ws, row, ncols, fill=HEAD_FILL, font=H2):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.fill = TITLE_FILL
    c.font = H1
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26


# ===========================================================================
def build():
    p = M.load_params()
    wb = Workbook()

    # -------------------- Лист «Допущения» --------------------
    ws = wb.active
    ws.title = "Допущения"
    title(ws, "ДОПУЩЕНИЯ (входные параметры — правьте синие значения)", 3)
    ws.cell(row=2, column=1, value="Параметр").font = BOLD
    ws.cell(row=2, column=2, value="Значение").font = BOLD
    ws.cell(row=2, column=3, value="Комментарий").font = BOLD
    style_header_row(ws, 2, 3)

    A = {}          # ключ -> адрес ячейки 'Допущения!Bxx'
    r = [3]

    def group(name):
        ws.cell(row=r[0], column=1, value=name).font = BOLD
        for c in (1, 2, 3):
            ws.cell(row=r[0], column=c).fill = GROUP_FILL
        r[0] += 1

    def inp(key, label, value, comment="", fmt=RUB):
        ws.cell(row=r[0], column=1, value=label)
        vc = ws.cell(row=r[0], column=2, value=value)
        vc.font = INPUT_FONT
        vc.number_format = fmt
        ws.cell(row=r[0], column=3, value=comment)
        A[key] = f"Допущения!$B${r[0]}"
        r[0] += 1

    group("Общие")
    inp("horizon", "Горизонт, лет", p["horizon_years"], "число лет прогноза", "0")
    inp("disc", "Ставка дисконтирования", p["discount_rate"], "ключевая + премия за риск", PCT)
    inp("infl", "Инфляция (рост цен/затрат)", p["inflation"], "годовая", PCT)
    inp("tax", "Налог ЕСХН", p["tax_eshn_rate"], "6% с прибыли", PCT)

    group("Земля")
    L = p["land"]
    inp("izhs_a", "Участок ИЖС, соток", L["izhs_area_sotka"], "под дом", "0")
    inp("izhs_p", "Цена ИЖС, ₽/сотка", L["izhs_price_per_sotka"], "Томский р-н 100–300 тыс")
    inp("farm_a", "Сельхозземля, га", L["farm_area_ha"], "поля/выпас", "0.0")
    inp("farm_p", "Цена сельхоз, ₽/га", L["farm_price_per_ha"], "до ~600 тыс/га")

    group("CAPEX (постройки/техника)")
    C = p["capex"]
    inp("house_m2", "Дом, м²", C["house_m2"], "", "0")
    inp("house_p", "Стройка, ₽/м²", C["house_price_per_m2"], "Томск, каркас/брус")
    inp("util", "Коммуникации, ₽", C["utilities"], "скважина, э/э, септик")
    inp("barn", "Коровник, ₽", C["barn"])
    inp("hay", "Сенник/склад, ₽", C["hay_storage"])
    inp("fence", "Ограждение/инфра, ₽", C["fencing_infra"])
    inp("mach", "Техника, ₽", C["machinery"])
    inp("green", "Теплицы/огород, ₽", C["greenhouses"])
    inp("reserve", "Резерв, % от CAPEX", C["reserve_pct"], "непредвиденное", PCT)

    group("Поголовье")
    H = p["herd"]
    inp("cows", "Коровы, голов", H["cows"], "", "0")
    inp("cow_p", "Цена коровы, ₽", H["cow_price"], "Томск 40–53 тыс")
    inp("young", "Молодняк, голов", H["young_stock"], "", "0")
    inp("young_p", "Цена молодняка, ₽", H["young_price"])
    inp("goats", "Козы, голов", H["goats"], "", "0")
    inp("goat_p", "Цена козы, ₽", H["goat_price"])
    inp("pigs", "Свиньи, голов", H["pigs"], "", "0")
    inp("pig_p", "Цена свиньи, ₽", H["pig_price"])
    inp("hens", "Куры, голов", H["hens"], "", "0")
    inp("hen_p", "Цена курицы, ₽", H["hen_price"])

    group("Продуктивность и цены")
    P = p["production"]
    inp("milk_y", "Надой, л/корову/год", P["milk_l_per_cow_year"], "", "0")
    inp("milk_p", "Цена молока, ₽/л", P["milk_price"], "КФХ-реализация")
    inp("cheese_s", "Доля молока в сыр", P["milk_to_cheese_share"], "", PCT)
    inp("cheese_y", "Выход сыра, кг/л", P["cheese_yield"], "", "0.00")
    inp("cheese_p", "Цена сыра, ₽/кг", P["cheese_price"])
    inp("meat_kg", "Мясо, кг/год", P["meat_kg_per_year"], "", "0")
    inp("meat_p", "Цена мяса, ₽/кг", P["meat_price"])
    inp("eggs", "Яиц на курицу/год", P["eggs_per_hen_year"], "", "0")
    inp("egg_p", "Цена яйца, ₽/шт", P["egg_price"])
    inp("veg", "Овощи, ₽/год", P["veg_revenue_year"])

    group("OPEX (годовые)")
    O = p["opex"]
    inp("feed_c", "Корм на корову, ₽/год", O["feed_purchased_per_cow"])
    inp("feed_off", "Экономия с 1 га, ₽", O["feed_self_offset_per_ha"], "свои поля")
    inp("vet", "Ветеринария на корову, ₽", O["vet_per_cow"])
    inp("ofeed", "Корм прочим, ₽/год", O["other_animals_feed"])
    inp("fuel", "ГСМ, ₽/год", O["fuel"])
    inp("elec", "Электричество, ₽/год", O["electricity"])
    inp("rep", "Ремонт/прочее, ₽/год", O["repairs_other"])
    inp("labor", "Наёмный труд, ₽/год", O["hired_labor"], "0 = семейный")

    group("Финансирование")
    F = p["financing"]
    inp("grant", "Грант КФХ, ₽", F["grant_amount"], "Агростартап Томск 3–6 млн")
    inp("cofin", "Софинансирование гранта, %", F["grant_own_cofinance_pct"], "мин. своих", PCT)
    inp("loan_s", "Доля кредита в CAPEX, %", F["loan_share_of_capex"], "", PCT)
    inp("loan_r", "Ставка кредита", F["loan_rate"], "льготная РСХБ до 5%", PCT)
    inp("loan_t", "Срок кредита, лет", F["loan_term_years"], "", "0")

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 34

    # -------------------- Лист CAPEX --------------------
    ws = wb.create_sheet("CAPEX")
    title(ws, "КАПИТАЛЬНЫЕ ВЛОЖЕНИЯ (формулы → лист «Допущения»)", 2)
    ws.cell(row=2, column=1, value="Статья").font = BOLD
    ws.cell(row=2, column=2, value="Сумма, ₽").font = BOLD
    style_header_row(ws, 2, 2)
    rows = [
        ("Земля ИЖС (участок под дом)", f"={A['izhs_a']}*{A['izhs_p']}"),
        ("Земля сельхоз (поля/выпас)", f"={A['farm_a']}*{A['farm_p']}"),
        ("Дом", f"={A['house_m2']}*{A['house_p']}"),
        ("Коммуникации", f"={A['util']}"),
        ("Коровник", f"={A['barn']}"),
        ("Сенник/склад", f"={A['hay']}"),
        ("Ограждение/инфраструктура", f"={A['fence']}"),
        ("Техника", f"={A['mach']}"),
        ("Теплицы/огород", f"={A['green']}"),
        ("Закупка скота (коровы)", f"={A['cows']}*{A['cow_p']}"),
        ("Закупка молодняка", f"={A['young']}*{A['young_p']}"),
        ("Прочие животные", f"={A['goats']}*{A['goat_p']}+{A['pigs']}*{A['pig_p']}+{A['hens']}*{A['hen_p']}"),
    ]
    rr = 3
    first = rr
    for label, f in rows:
        ws.cell(row=rr, column=1, value=label)
        c = ws.cell(row=rr, column=2, value=f); c.number_format = RUB
        rr += 1
    last = rr - 1
    ws.cell(row=rr, column=1, value="Подытог").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"=SUM(B{first}:B{last})"); c.number_format = RUB; c.font = BOLD
    sub = rr; rr += 1
    ws.cell(row=rr, column=1, value="Резерв").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"=B{sub}*{A['reserve']}"); c.number_format = RUB; c.font = BOLD
    res = rr; rr += 1
    ws.cell(row=rr, column=1, value="ИТОГО CAPEX").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"=B{sub}+B{res}"); c.number_format = RUB; c.font = BOLD
    c.fill = TOTAL_FILL
    ws.cell(row=rr, column=1).fill = TOTAL_FILL
    CAPEX_TOTAL = f"CAPEX!$B${rr}"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16

    # -------------------- Лист P&L --------------------
    ws = wb.create_sheet("P&L")
    title(ws, "ГОДОВОЙ P&L: ВЫРУЧКА − OPEX = EBITDA", 2)
    rr = 2
    ws.cell(row=rr, column=1, value="ВЫРУЧКА").font = BOLD
    ws.cell(row=rr, column=1).fill = GROUP_FILL; ws.cell(row=rr, column=2).fill = GROUP_FILL
    rr += 1
    rev_rows = [
        ("Молоко (сырое)", f"={A['cows']}*{A['milk_y']}*(1-{A['cheese_s']})*{A['milk_p']}"),
        ("Сыр/творог (переработка)", f"={A['cows']}*{A['milk_y']}*{A['cheese_s']}*{A['cheese_y']}*{A['cheese_p']}"),
        ("Мясо", f"={A['meat_kg']}*{A['meat_p']}"),
        ("Яйца", f"={A['hens']}*{A['eggs']}*{A['egg_p']}"),
        ("Овощи/огород", f"={A['veg']}"),
    ]
    rfirst = rr
    for label, f in rev_rows:
        ws.cell(row=rr, column=1, value=label)
        c = ws.cell(row=rr, column=2, value=f); c.number_format = RUB
        rr += 1
    rlast = rr - 1
    ws.cell(row=rr, column=1, value="ИТОГО выручка").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"=SUM(B{rfirst}:B{rlast})"); c.number_format = RUB; c.font = BOLD
    REV = f"'P&L'!$B${rr}"; rr += 2

    ws.cell(row=rr, column=1, value="OPEX").font = BOLD
    ws.cell(row=rr, column=1).fill = GROUP_FILL; ws.cell(row=rr, column=2).fill = GROUP_FILL
    rr += 1
    opex_rows = [
        ("Корма (за вычетом своих полей)", f"=MAX(0,{A['cows']}*{A['feed_c']}-{A['farm_a']}*{A['feed_off']})"),
        ("Корм прочим животным", f"={A['ofeed']}"),
        ("Ветеринария/осеменение", f"={A['cows']}*{A['vet']}"),
        ("ГСМ", f"={A['fuel']}"),
        ("Электричество", f"={A['elec']}"),
        ("Ремонт/прочее", f"={A['rep']}"),
        ("Наёмный труд", f"={A['labor']}"),
    ]
    ofirst = rr
    for label, f in opex_rows:
        ws.cell(row=rr, column=1, value=label)
        c = ws.cell(row=rr, column=2, value=f); c.number_format = RUB
        rr += 1
    olast = rr - 1
    ws.cell(row=rr, column=1, value="ИТОГО OPEX").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"=SUM(B{ofirst}:B{olast})"); c.number_format = RUB; c.font = BOLD
    OPEX = f"'P&L'!$B${rr}"; rr += 1
    ws.cell(row=rr, column=1, value="EBITDA (выручка − OPEX)").font = BOLD
    c = ws.cell(row=rr, column=2, value=f"={REV}-{OPEX}"); c.number_format = RUB; c.font = BOLD
    c.fill = TOTAL_FILL; ws.cell(row=rr, column=1).fill = TOTAL_FILL
    EBITDA = f"'P&L'!$B${rr}"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16

    # -------------------- Лист Юнит-экономика --------------------
    ws = wb.create_sheet("Юнит-экономика")
    title(ws, "ЮНИТ-ЭКОНОМИКА", 2)
    ws.cell(row=2, column=1, value="Показатель").font = BOLD
    ws.cell(row=2, column=2, value="Значение").font = BOLD
    style_header_row(ws, 2, 2)
    u_rows = [
        ("Выручка с 1 коровы (молоко), ₽/год", f"={A['milk_y']}*{A['milk_p']}"),
        ("Прямые затраты на 1 корову, ₽/год", f"={A['feed_c']}+{A['vet']}"),
        ("Маржа на 1 корову, ₽/год", f"={A['milk_y']}*{A['milk_p']}-({A['feed_c']}+{A['vet']})"),
        ("Эффект 1 га кормовых, ₽/год", f"={A['feed_off']}"),
        ("Безубыточное стадо vs пост. затраты, голов",
         f"=({A['ofeed']}+{A['fuel']}+{A['elec']}+{A['rep']}+{A['labor']})/"
         f"({A['milk_y']}*{A['milk_p']}-{A['feed_c']}-{A['vet']})"),
    ]
    rr = 3
    for label, f in u_rows:
        ws.cell(row=rr, column=1, value=label)
        fmt = "0.0" if "голов" in label else RUB
        c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
        rr += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16

    # -------------------- Лист Денежный поток (база) --------------------
    ws = wb.create_sheet("Денежный поток (база)")
    title(ws, "ДЕНЕЖНЫЙ ПОТОК — БАЗОВЫЙ ПРОГНОЗ (формулы)", 4)
    ws.cell(row=2, column=1, value="Год")
    cols = list(M.FIN_SCENARIOS.values())
    for i, name in enumerate(cols):
        ws.cell(row=2, column=2 + i, value=name)
    style_header_row(ws, 2, 1 + len(cols))

    # вспомогательные именованные адреса
    n = p["horizon_years"]
    # столбцы: B=own, C=grant, D=loan
    # year 0 — собственные вложения
    own0 = {
        "own": f"=-{CAPEX_TOTAL}",
        "grant": f"=-MAX({CAPEX_TOTAL}-{A['grant']},{CAPEX_TOTAL}*{A['cofin']})",
        "loan": f"=-{CAPEX_TOTAL}*(1-{A['loan_s']})",
    }
    keys = list(M.FIN_SCENARIOS.keys())
    rr = 3
    ws.cell(row=rr, column=1, value=0)
    for i, k in enumerate(keys):
        c = ws.cell(row=rr, column=2 + i, value=own0[k]); c.number_format = RUB
    rr += 1
    # платёж по кредиту (annuity) — вынесем в ячейку справа
    pmt_cell = "G3"
    ws.cell(row=3, column=6, value="Платёж по кредиту, ₽/год:").font = BOLD
    pc = ws.cell(row=3, column=7,
                 value=f"={CAPEX_TOTAL}*{A['loan_s']}*{A['loan_r']}/(1-(1+{A['loan_r']})^-{A['loan_t']})")
    pc.number_format = RUB
    for t in range(1, n + 1):
        ws.cell(row=rr, column=1, value=t)
        # EBITDA с инфляцией, минус налог
        ebt = f"({EBITDA}*(1+{A['infl']})^{t-1})"
        net = f"={ebt}*(1-{A['tax']})"
        # own
        c = ws.cell(row=rr, column=2, value=net); c.number_format = RUB
        # grant (та же операционка)
        c = ws.cell(row=rr, column=3, value=net); c.number_format = RUB
        # loan: минус платёж, пока t<=срок
        loan_net = f"={ebt}*(1-{A['tax']})-IF({t}<={A['loan_t']},${pmt_cell},0)"
        c = ws.cell(row=rr, column=4, value=loan_net); c.number_format = RUB
        rr += 1

    last_year_row = rr - 1
    # метрики
    rr += 1
    metrics_start = rr
    ws.cell(row=rr, column=1, value="NPV").font = BOLD
    for i, k in enumerate(keys):
        col = get_column_letter(2 + i)
        # NPV = поток года0 + NPV(rate, потоки 1..n)
        f = f"={col}3+NPV({A['disc']},{col}4:{col}{last_year_row})"
        c = ws.cell(row=rr, column=2 + i, value=f); c.number_format = RUB; c.font = BOLD
    rr += 1
    ws.cell(row=rr, column=1, value="IRR").font = BOLD
    for i, k in enumerate(keys):
        col = get_column_letter(2 + i)
        f = f"=IFERROR(IRR({col}3:{col}{last_year_row}),\"—\")"
        c = ws.cell(row=rr, column=2 + i, value=f); c.number_format = PCT; c.font = BOLD
    rr += 1
    ws.cell(row=rr, column=1, value="Свои вложения (год 0), ₽").font = BOLD
    for i, k in enumerate(keys):
        col = get_column_letter(2 + i)
        c = ws.cell(row=rr, column=2 + i, value=f"=-{col}3"); c.number_format = RUB; c.font = BOLD
    ws.column_dimensions["A"].width = 26
    for col in "BCD":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 14

    # -------------------- Лист Прогноз (3 сценария) --------------------
    res = M.run_all(p)
    ws = wb.create_sheet("Прогноз (3 сценария)")
    title(ws, "ПРОГНОЗЫ: ПЕССИМИСТИЧНЫЙ / БАЗОВЫЙ / ОПТИМИСТИЧНЫЙ", 5)
    # блок множителей
    rr = 2
    ws.cell(row=rr, column=1, value="Множители сценариев").font = BOLD
    for c in range(1, 6):
        ws.cell(row=rr, column=c).fill = GROUP_FILL
    rr += 1
    mult_keys = [("milk_price", "Цена молока"), ("milk_yield", "Надой"),
                 ("feed_cost", "Стоимость кормов"), ("capex", "CAPEX"),
                 ("veg_revenue", "Выручка овощи"), ("meat_price", "Цена мяса"),
                 ("discount_add", "Δ ставки дисконт.")]
    fnames = list(p["forecasts"].keys())
    ws.cell(row=rr, column=1, value="Драйвер").font = BOLD
    for i, fn in enumerate(fnames):
        ws.cell(row=rr, column=2 + i, value=fn).font = BOLD
    style_header_row(ws, rr, 1 + len(fnames), fill=HEAD_FILL)
    rr += 1
    for mk, mlabel in mult_keys:
        ws.cell(row=rr, column=1, value=mlabel)
        for i, fn in enumerate(fnames):
            v = p["forecasts"][fn][mk]
            cell = ws.cell(row=rr, column=2 + i, value=v)
            cell.number_format = "0%" if mk != "discount_add" else "+0%;-0%"
        rr += 1
    rr += 1

    # сводные показатели
    ws.cell(row=rr, column=1, value="Показатель").font = BOLD
    for i, fn in enumerate(fnames):
        ws.cell(row=rr, column=2 + i, value=fn).font = BOLD
    style_header_row(ws, rr, 1 + len(fnames))
    rr += 1

    def row_vals(label, getter, fmt=RUB):
        ws.cell(row=rr[0], column=1, value=label)
        for i, fn in enumerate(fnames):
            c = ws.cell(row=rr[0], column=2 + i, value=getter(res[fn]))
            c.number_format = fmt
        rr[0] += 1

    rr = [rr]
    row_vals("CAPEX, ₽", lambda d: round(d["capex_total"]))
    row_vals("Выручка, ₽/год", lambda d: round(d["annual_rev"]))
    row_vals("OPEX, ₽/год", lambda d: round(d["annual_opex"]))
    row_vals("EBITDA, ₽/год", lambda d: round(d["ebitda"]))
    rr = rr[0]
    # по каждой схеме финансирования — окупаемость / NPV / IRR
    for k, kname in M.FIN_SCENARIOS.items():
        ws.cell(row=rr, column=1, value=kname).font = BOLD
        for c in range(1, 2 + len(fnames)):
            ws.cell(row=rr, column=c).fill = GROUP_FILL
        rr += 1
        for metric, fmt, fn_get in [
            ("Окупаемость, лет", "0.0", lambda m: m["payback"] if m["payback"] is not None else "—"),
            ("NPV, ₽", RUB, lambda m: round(m["npv"])),
            ("IRR", PCT, lambda m: m["irr"] if m["irr"] is not None else "—"),
        ]:
            ws.cell(row=rr, column=1, value="   " + metric)
            for i, fn in enumerate(fnames):
                m = res[fn]["fin"][k]
                val = fn_get(m)
                c = ws.cell(row=rr, column=2 + i, value=val)
                if isinstance(val, (int, float)):
                    c.number_format = fmt
                    if metric.startswith("NPV"):
                        c.fill = GOOD_FILL if val >= 0 else BAD_FILL
            rr += 1
    ws.column_dimensions["A"].width = 30
    for col in "BCD":
        ws.column_dimensions[col].width = 17

    # -------------------- Лист Чувствительность --------------------
    ws = wb.create_sheet("Чувствительность")
    title(ws, "ЧУВСТВИТЕЛЬНОСТЬ NPV (схема «Грант КФХ», базовый прогноз)", 4)
    ws.cell(row=2, column=1,
            value="Строки — цена молока, столбцы — стоимость кормов. Значения: NPV, ₽.")
    base = M.apply_forecast(p, "Базовый")
    base_milk = base["production"]["milk_price"]
    base_feed = base["opex"]["feed_purchased_per_cow"]
    feed_deltas = [-0.2, -0.1, 0.0, 0.1, 0.2]
    milk_deltas = [-0.2, -0.1, 0.0, 0.1, 0.2]
    hr = 4
    ws.cell(row=hr, column=1, value="Молоко \\ Корма").font = BOLD
    for j, fd in enumerate(feed_deltas):
        ws.cell(row=hr, column=2 + j, value=f"{fd:+.0%}").font = BOLD
    style_header_row(ws, hr, 1 + len(feed_deltas))
    rr = hr + 1
    for md in milk_deltas:
        ws.cell(row=rr, column=1, value=f"{md:+.0%}").font = BOLD
        ws.cell(row=rr, column=1).fill = GROUP_FILL
        for j, fd in enumerate(feed_deltas):
            pp = M.apply_forecast(p, "Базовый")
            pp["production"]["milk_price"] = base_milk * (1 + md)
            pp["opex"]["feed_purchased_per_cow"] = base_feed * (1 + fd)
            _, rv = M.calc_revenue(pp)
            _, ox = M.calc_opex(pp)
            _, ct = M.calc_capex(pp)
            m = M.scenario_metrics(pp, ct, rv, ox, "grant")
            c = ws.cell(row=rr, column=2 + j, value=round(m["npv"]))
            c.number_format = RUB
            c.fill = GOOD_FILL if m["npv"] >= 0 else BAD_FILL
        rr += 1
    ws.column_dimensions["A"].width = 16
    for j in range(len(feed_deltas)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 15

    # -------------------- Лист Источники --------------------
    ws = wb.create_sheet("Источники")
    title(ws, "ИСТОЧНИКИ ДАННЫХ (Томск, 2026)", 2)
    src = [
        ("Участки ИЖС Томский р-н (100–300 тыс/сотка)", "tomsk.cian.ru / tomsk.sibdom.ru / m2.ru"),
        ("Сельхозземля (до ~600 тыс/га)", "tomsk.move.ru / tomsk.cian.ru"),
        ("Строительство дома (Томск)", "dom-stroy70.ru / tomsk.kamprok.ru"),
        ("Закупочная цена молока (выс. сорт ~40 ₽/л, КФХ меньше)", "vtomske.ru / souzmoloko.ru"),
        ("Цена КРС (дойная 40–53 тыс ₽)", "sobut.ru / agroserver.ru"),
        ("Грант «Агростартап» Томск (3–6 млн, до 7 на КРС)", "svetich.info / svoefermerstvo.ru"),
        ("Льготный кредит АПК (ставка до 5%, срок до 7 лет)", "rshb.ru / svoefermerstvo.ru"),
    ]
    ws.cell(row=2, column=1, value="Параметр").font = BOLD
    ws.cell(row=2, column=2, value="Источник").font = BOLD
    style_header_row(ws, 2, 2)
    rr = 3
    for a, b in src:
        ws.cell(row=rr, column=1, value=a)
        ws.cell(row=rr, column=2, value=b)
        rr += 1
    ws.cell(row=rr + 1, column=1,
            value="Цены ориентировочные (диапазоны рынка). Перед решением проверьте по конкретному участку.")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 40

    wb.save(OUT)
    print("[OK] Excel сохранён:", OUT)


if __name__ == "__main__":
    build()
